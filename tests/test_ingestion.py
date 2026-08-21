"""
Unit tests for Module 1 - Document Ingestion.

Covers ingestion/schema.py, ingestion/loaders.py, and
ingestion/dedupe.py. PDF/DOCX test fixtures are generated on the fly
with PyMuPDF and python-docx respectively, so these tests have no
external file dependencies.
"""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

import fitz  # PyMuPDF
import pytest
from docx import Document as DocxDocument

from ingestion.dedupe import compute_checksum, dedupe_by_checksum
from ingestion.loaders import (
    DocumentLoader,
    DocumentLoaderError,
    UnsupportedFormatError,
    extract_docx,
    extract_html,
    extract_pdf,
)
from ingestion.schema import (
    DocumentFormat,
    RawDocument,
    SchemaValidationError,
    SourceSystem,
)


# --------------------------------------------------------------------------
# Fixtures
# --------------------------------------------------------------------------


@pytest.fixture
def sample_pdf(tmp_path: Path) -> Path:
    """Creates a minimal one-page PDF with known text content."""
    pdf_path = tmp_path / "sample.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Employees are entitled to 12 casual leave days.")
    doc.save(pdf_path)
    doc.close()
    return pdf_path


@pytest.fixture
def empty_pdf(tmp_path: Path) -> Path:
    """Creates a PDF with a blank page (no text) to exercise the failure path."""
    pdf_path = tmp_path / "empty.pdf"
    doc = fitz.open()
    doc.new_page()
    doc.save(pdf_path)
    doc.close()
    return pdf_path


@pytest.fixture
def sample_docx(tmp_path: Path) -> Path:
    """Creates a minimal DOCX with a paragraph and a table."""
    docx_path = tmp_path / "sample.docx"
    document = DocxDocument()
    document.add_paragraph("VPN access requires manager approval.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Department"
    table.rows[0].cells[1].text = "Finance"
    document.save(str(docx_path))
    return docx_path


@pytest.fixture
def sample_html(tmp_path: Path) -> Path:
    """Creates a minimal HTML file with a script tag that should be stripped."""
    html_path = tmp_path / "sample.html"
    html_path.write_text(
        "<html><head><style>body{color:red;}</style></head>"
        "<body><h1>Leave Policy</h1>"
        "<p>Casual leave is 12 days per year.</p>"
        "<script>console.log('should not appear');</script>"
        "</body></html>",
        encoding="utf-8",
    )
    return html_path


def _make_raw_document(text: str = "hello world", checksum: str | None = None) -> RawDocument:
    return RawDocument(
        document_id="pdf_abc123",
        text=text,
        source_path="/data/raw/sample.pdf",
        source_system=SourceSystem.LOCAL_FILESYSTEM,
        document_format=DocumentFormat.PDF,
        collected_at=datetime.now(timezone.utc),
        checksum=checksum or compute_checksum(text.encode("utf-8")),
    )


# --------------------------------------------------------------------------
# schema.py
# --------------------------------------------------------------------------


class TestRawDocument:
    def test_char_count_and_file_name(self) -> None:
        doc = _make_raw_document(text="abcde")
        assert doc.char_count == 5
        assert doc.file_name == "sample.pdf"

    def test_missing_checksum_raises(self) -> None:
        with pytest.raises(SchemaValidationError):
            RawDocument(
                document_id="pdf_x",
                text="text",
                source_path="/x.pdf",
                source_system=SourceSystem.LOCAL_FILESYSTEM,
                document_format=DocumentFormat.PDF,
                collected_at=datetime.now(timezone.utc),
                checksum="",
            )

    def test_empty_document_id_raises(self) -> None:
        with pytest.raises(SchemaValidationError):
            RawDocument(
                document_id="",
                text="text",
                source_path="/x.pdf",
                source_system=SourceSystem.LOCAL_FILESYSTEM,
                document_format=DocumentFormat.PDF,
                collected_at=datetime.now(timezone.utc),
                checksum="deadbeef",
            )

    def test_empty_text_logs_warning_but_does_not_raise(self, caplog) -> None:
        with caplog.at_level("WARNING"):
            doc = _make_raw_document(text="   ")
        assert doc.text.strip() == ""
        assert any("empty or whitespace" in record.message for record in caplog.records)

    def test_is_immutable(self) -> None:
        doc = _make_raw_document()
        with pytest.raises(AttributeError):
            doc.text = "mutated"  # type: ignore[misc]

    def test_to_metadata_dict_contains_core_fields(self) -> None:
        doc = _make_raw_document()
        metadata = doc.to_metadata_dict()
        assert metadata["document_id"] == doc.document_id
        assert metadata["checksum"] == doc.checksum
        assert metadata["source_system"] == "local_filesystem"


# --------------------------------------------------------------------------
# dedupe.py
# --------------------------------------------------------------------------


class TestComputeChecksum:
    def test_same_content_same_checksum(self) -> None:
        content = b"identical bytes"
        assert compute_checksum(content) == compute_checksum(content)

    def test_different_content_different_checksum(self) -> None:
        assert compute_checksum(b"a") != compute_checksum(b"b")

    def test_empty_content_raises(self) -> None:
        with pytest.raises(ValueError):
            compute_checksum(b"")

    def test_unsupported_algorithm_raises(self) -> None:
        with pytest.raises(ValueError):
            compute_checksum(b"data", algorithm="sha3000")

    def test_md5_algorithm_supported(self) -> None:
        digest = compute_checksum(b"data", algorithm="md5")
        assert len(digest) == 32  # MD5 hex digest length


class TestDedupeByChecksum:
    def test_removes_exact_duplicates(self) -> None:
        doc_a = _make_raw_document(text="same content", checksum="checksum-1")
        doc_b = _make_raw_document(text="same content", checksum="checksum-1")
        doc_c = _make_raw_document(text="different content", checksum="checksum-2")

        result = dedupe_by_checksum([doc_a, doc_b, doc_c])

        assert len(result) == 2
        assert result[0] is doc_a
        assert result[1] is doc_c

    def test_no_duplicates_returns_all(self) -> None:
        docs = [
            _make_raw_document(checksum="c1"),
            _make_raw_document(checksum="c2"),
            _make_raw_document(checksum="c3"),
        ]
        result = dedupe_by_checksum(docs)
        assert len(result) == 3

    def test_empty_list_returns_empty_list(self) -> None:
        assert dedupe_by_checksum([]) == []


# --------------------------------------------------------------------------
# loaders.py - extraction functions
# --------------------------------------------------------------------------


class TestExtractPdf:
    def test_extracts_known_text(self, sample_pdf: Path) -> None:
        text = extract_pdf(sample_pdf)
        assert "12 casual leave days" in text

    def test_empty_pdf_raises(self, empty_pdf: Path) -> None:
        with pytest.raises(DocumentLoaderError):
            extract_pdf(empty_pdf)

    def test_missing_file_raises(self, tmp_path: Path) -> None:
        with pytest.raises(DocumentLoaderError):
            extract_pdf(tmp_path / "does_not_exist.pdf")


class TestExtractDocx:
    def test_extracts_paragraph_and_table_text(self, sample_docx: Path) -> None:
        text = extract_docx(sample_docx)
        assert "VPN access requires manager approval." in text
        assert "Finance" in text

    def test_missing_file_raises(self, tmp_path: Path) -> None:
        with pytest.raises(DocumentLoaderError):
            extract_docx(tmp_path / "does_not_exist.docx")


class TestExtractHtml:
    def test_extracts_visible_text_only(self, sample_html: Path) -> None:
        text = extract_html(sample_html)
        assert "Casual leave is 12 days per year." in text
        assert "should not appear" not in text
        assert "color:red" not in text

    def test_missing_file_raises(self, tmp_path: Path) -> None:
        with pytest.raises(DocumentLoaderError):
            extract_html(tmp_path / "does_not_exist.html")


# --------------------------------------------------------------------------
# loaders.py - DocumentLoader
# --------------------------------------------------------------------------


class TestDocumentLoader:
    def test_supports(self, sample_pdf: Path) -> None:
        loader = DocumentLoader()
        assert loader.supports(sample_pdf) is True
        assert loader.supports(Path("file.xyz")) is False

    def test_load_pdf_returns_populated_raw_document(self, sample_pdf: Path) -> None:
        loader = DocumentLoader(source_system=SourceSystem.LOCAL_FILESYSTEM)
        doc = loader.load(sample_pdf)

        assert isinstance(doc, RawDocument)
        assert doc.document_format == DocumentFormat.PDF
        assert doc.source_system == SourceSystem.LOCAL_FILESYSTEM
        assert doc.checksum
        assert "12 casual leave days" in doc.text

    def test_load_docx_returns_populated_raw_document(self, sample_docx: Path) -> None:
        loader = DocumentLoader()
        doc = loader.load(sample_docx)
        assert doc.document_format == DocumentFormat.DOCX
        assert "VPN access" in doc.text

    def test_load_html_returns_populated_raw_document(self, sample_html: Path) -> None:
        loader = DocumentLoader()
        doc = loader.load(sample_html)
        assert doc.document_format == DocumentFormat.HTML
        assert "Casual leave" in doc.text

    def test_load_missing_file_raises_file_not_found(self, tmp_path: Path) -> None:
        loader = DocumentLoader()
        with pytest.raises(FileNotFoundError):
            loader.load(tmp_path / "missing.pdf")

    def test_load_unsupported_extension_raises(self, tmp_path: Path) -> None:
        unsupported = tmp_path / "notes.txt"
        unsupported.write_text("plain text", encoding="utf-8")
        loader = DocumentLoader()
        with pytest.raises(UnsupportedFormatError):
            loader.load(unsupported)

    def test_two_identical_files_produce_same_document_id(
        self, sample_pdf: Path, tmp_path: Path
    ) -> None:
        loader = DocumentLoader()
        copy_path = tmp_path / "copy.pdf"
        copy_path.write_bytes(sample_pdf.read_bytes())

        doc_a = loader.load(sample_pdf)
        doc_b = loader.load(copy_path)

        assert doc_a.checksum == doc_b.checksum
        assert doc_a.document_id == doc_b.document_id

    def test_load_directory_skips_unsupported_and_loads_rest(
        self, tmp_path: Path, sample_pdf: Path, sample_docx: Path, sample_html: Path
    ) -> None:
        (tmp_path / "ignore_me.xyz").write_text("irrelevant", encoding="utf-8")
        loader = DocumentLoader()

        documents = loader.load_directory(tmp_path, recursive=False)

        formats = {doc.document_format for doc in documents}
        assert formats == {DocumentFormat.PDF, DocumentFormat.DOCX, DocumentFormat.HTML}
        assert len(documents) == 3

    def test_load_directory_on_file_raises_not_a_directory(self, sample_pdf: Path) -> None:
        loader = DocumentLoader()
        with pytest.raises(NotADirectoryError):
            loader.load_directory(sample_pdf)

    def test_load_directory_recursive_finds_nested_files(
        self, tmp_path: Path, sample_pdf: Path
    ) -> None:
        # Use an isolated root so this test's fixture directory doesn't also
        # contain the sample_pdf fixture's own copy (both fixtures share
        # tmp_path by default).
        root = tmp_path / "root"
        nested_dir = root / "nested"
        nested_dir.mkdir(parents=True)
        nested_pdf = nested_dir / "nested.pdf"
        nested_pdf.write_bytes(sample_pdf.read_bytes())

        loader = DocumentLoader()
        documents = loader.load_directory(root, recursive=True)

        assert len(documents) == 1
        assert documents[0].source_path.endswith("nested.pdf")

    def test_load_directory_non_recursive_ignores_nested_files(
        self, tmp_path: Path, sample_pdf: Path
    ) -> None:
        root = tmp_path / "root2"
        nested_dir = root / "nested"
        nested_dir.mkdir(parents=True)
        (nested_dir / "nested.pdf").write_bytes(sample_pdf.read_bytes())
        (root / "top_level.pdf").write_bytes(sample_pdf.read_bytes())

        loader = DocumentLoader()
        documents = loader.load_directory(root, recursive=False)

        assert len(documents) == 1
        assert documents[0].source_path.endswith("top_level.pdf")
